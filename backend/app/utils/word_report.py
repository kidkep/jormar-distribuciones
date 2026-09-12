import io
import os

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

LOGO_PATH = os.path.join(os.path.dirname(__file__), "logo.png")

GOLD = "C49828"
GOLD_DARK = "967019"
DARK = "282828"
GRAY = "6B7280"
LIGHT = "F5F3EC"
MID_GRAY = "E2E0D8"
GREEN = "168A34"
RED = "DC2626"
ORANGE = "C2410C"


def _set_cell_shading(cell, color_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), color_hex)
    tcPr.append(shd)


def _set_cell_text(cell, text, bold=False, color=DARK, size=9.5, align="left"):
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = {
        "left": WD_ALIGN_PARAGRAPH.LEFT,
        "right": WD_ALIGN_PARAGRAPH.RIGHT,
        "center": WD_ALIGN_PARAGRAPH.CENTER,
    }.get(align, WD_ALIGN_PARAGRAPH.LEFT)
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    run = p.add_run(str(text))
    run.font.name = "Calibri"
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)


def _paragraph_border(p, color=GOLD, size=12, space=4, side="bottom"):
    pPr = p._p.get_or_add_pPr()
    pBdr = OxmlElement("w:pBdr")
    border = OxmlElement(f"w:{side}")
    border.set(qn("w:val"), "single")
    border.set(qn("w:sz"), str(size))
    border.set(qn("w:space"), str(space))
    border.set(qn("w:color"), color)
    pBdr.append(border)
    pPr.append(pBdr)


def _add_section_heading(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(14)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.keep_with_next = True
    run = p.add_run(text)
    run.font.name = "Calibri"
    run.font.size = Pt(14)
    run.font.bold = True
    run.font.color.rgb = RGBColor.from_string(GOLD_DARK)
    _paragraph_border(p)
    return p


def _add_paragraph(doc, text, size=10.5, bold=False, color=DARK, italic=False, space_after=6, align=None):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(space_after)
    if align:
        p.alignment = align
    run = p.add_run(text)
    run.font.name = "Calibri"
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = RGBColor.from_string(color)
    return p


def _add_bullet(doc, label, text, label_color=GOLD_DARK, size=10):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(4)
    if label:
        run_label = p.add_run(f"{label}: ")
        run_label.font.name = "Calibri"
        run_label.font.size = Pt(size)
        run_label.font.bold = True
        run_label.font.color.rgb = RGBColor.from_string(label_color)
    run = p.add_run(text)
    run.font.name = "Calibri"
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(DARK)
    return p


def _add_table(doc, headers, rows, col_widths_cm=None, right_cols=None, last_bold=False):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    if col_widths_cm:
        for i, w in enumerate(col_widths_cm):
            for row in table.rows:
                row.cells[i].width = Cm(w)

    hdr = table.rows[0]
    for i, h in enumerate(headers):
        _set_cell_shading(hdr.cells[i], GOLD_DARK)
        _set_cell_text(hdr.cells[i], h, bold=True, color="FFFFFF", align="center")

    right_cols = right_cols or []
    for idx, r in enumerate(rows):
        cells = table.add_row().cells
        row_bold = last_bold and idx == len(rows) - 1
        for i, value in enumerate(r):
            _set_cell_shading(cells[i], LIGHT if len(table.rows) % 2 == 0 else "FFFFFF")
            _set_cell_text(
                cells[i], value,
                bold=row_bold,
                align="right" if i in right_cols else "left",
            )

    for i, w in enumerate(col_widths_cm or []):
        for row in table.rows:
            row.cells[i].width = Cm(w)
    return table


def _add_footer_numbering(doc):
    section = doc.sections[0]
    footer = section.footer
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("JORMAR DISTRIBUCIONES - Informe de gestion  |  Pagina ")
    run.font.name = "Calibri"
    run.font.size = Pt(8)
    run.font.color.rgb = RGBColor.from_string(GRAY)

    fldChar1 = OxmlElement("w:fldChar")
    fldChar1.set(qn("w:fldCharType"), "begin")
    instrText = OxmlElement("w:instrText")
    instrText.set(qn("xml:space"), "preserve")
    instrText.text = "PAGE"
    fldChar2 = OxmlElement("w:fldChar")
    fldChar2.set(qn("w:fldCharType"), "end")

    run2 = p.add_run()
    run2.font.name = "Calibri"
    run2.font.size = Pt(8)
    run2.font.color.rgb = RGBColor.from_string(GRAY)
    run2._r.append(fldChar1)
    run2._r.append(instrText)
    run2._r.append(fldChar2)


def _money(v) -> str:
    return f"$ {float(v or 0):,.0f}"


def _num(v) -> str:
    v = float(v or 0)
    if v == int(v):
        return f"{int(v):,}"
    return f"{v:,.1f}"


def _pct_color(value):
    return GREEN if value >= 0 else RED


def generate_report_docx_bytes(data: dict) -> bytes:
    doc = Document()

    for section in doc.sections:
        section.top_margin = Cm(2)
        section.bottom_margin = Cm(2)
        section.left_margin = Cm(2)
        section.right_margin = Cm(2)

    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(10.5)
    normal.paragraph_format.space_after = Pt(6)

    _add_footer_numbering(doc)

    # ---------- CABECERA ----------
    if os.path.exists(LOGO_PATH):
        p_logo = doc.add_paragraph()
        p_logo.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_logo.paragraph_format.space_after = Pt(2)
        run_logo = p_logo.add_run()
        run_logo.add_picture(LOGO_PATH, width=Cm(4.2))

    t = doc.add_paragraph()
    t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    t.paragraph_format.space_after = Pt(1)
    run_t = t.add_run("JORMAR DISTRIBUCIONES")
    run_t.font.name = "Calibri"
    run_t.font.size = Pt(22)
    run_t.font.bold = True
    run_t.font.color.rgb = RGBColor.from_string(GOLD_DARK)

    st = doc.add_paragraph()
    st.alignment = WD_ALIGN_PARAGRAPH.CENTER
    st.paragraph_format.space_after = Pt(1)
    run_st = st.add_run("INFORME EJECUTIVO DE GESTIÓN")
    run_st.font.name = "Calibri"
    run_st.font.size = Pt(14)
    run_st.font.bold = True
    run_st.font.color.rgb = RGBColor.from_string(GOLD)

    _paragraph_border(st, color=GOLD, size=16, side="bottom")

    info = doc.add_paragraph()
    info.alignment = WD_ALIGN_PARAGRAPH.CENTER
    info.paragraph_format.space_after = Pt(10)
    run_info = info.add_run(
        f"NIT 931814237 - Mariquita, Tolima   |   "
        f"Periodo: {data['periodo']['inicio_legible']} al {data['periodo']['fin_legible']}   |   "
        f"Generado: {data['generado']}"
    )
    run_info.font.name = "Calibri"
    run_info.font.size = Pt(9)
    run_info.font.color.rgb = RGBColor.from_string(GRAY)

    # ---------- SALUD DE LA EMPRESA ----------
    salud_t = doc.add_table(rows=1, cols=1)
    salud_t.style = "Table Grid"
    salud_t.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = salud_t.cell(0, 0)
    _set_cell_shading(cell, LIGHT)
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r1 = p.add_run("SALUD DE LA EMPRESA: ")
    r1.font.name = "Calibri"
    r1.font.size = Pt(12)
    r1.font.bold = True
    r1.font.color.rgb = RGBColor.from_string(GOLD_DARK)
    r2 = p.add_run(data["salud"].upper())
    r2.font.name = "Calibri"
    r2.font.size = Pt(12)
    r2.font.bold = True
    color_salud = GREEN if data["salud"] == "Saludable" else (ORANGE if data["salud"] == "Con atencion" else RED)
    r2.font.color.rgb = RGBColor.from_string(color_salud)

    # ---------- RESUMEN EJECUTIVO ----------
    _add_section_heading(doc, "1. RESUMEN EJECUTIVO")
    _add_paragraph(doc, data["resumen"])

    # ---------- INDICADORES CLAVE ----------
    _add_section_heading(doc, "2. INDICADORES CLAVE")
    v = data["ventas"]
    g = data["gastos"]
    r = data["rentabilidad"]
    cmp = data["comparativo"]
    kpi_rows = [
        ("Facturado del periodo (todo lo vendido)", _money(v["facturado"])),
        ("Dinero real recibido (contado + abonos)", _money(v["recibido_real"])),
        ("Gastos del periodo", _money(g["total"])),
        ("Ganancia de caja (recibido real - gastos)", _money(r["ganancia_caja"])),
        ("Utilidad comercial estimada (menos costo mercancia)", _money(r["utilidad_comercial"])),
        ("Crecimiento de facturado vs periodo anterior", f"{_num(cmp['crecimiento_facturado'])}%"),
        ("Ventas totales en el periodo", f"{_num(v['cantidad'])}"),
        ("Ticket promedio", _money(v["ticket_promedio"])),
        ("Inventario a costo de compra", _money(data["inventario"]["costo"])),
        ("Cartera por cobrar (creditos pendientes)", _money(data["cartera"]["deuda_pendiente"])),
        ("Clientes activos", f"{_num(data['clientes']['total'])}"),
    ]
    _add_table(doc, ["Indicador", "Valor"], kpi_rows, col_widths_cm=[13.5, 3.5], right_cols=[1])

    # ---------- VENTAS ----------
    _add_section_heading(doc, "3. VENTAS")
    venta_rows = [
        ("Facturado total", _money(v["facturado"])),
        ("Ventas de contado", _money(v["contado"])),
        ("Ventas a credito", _money(v["credito_vendido"])),
        ("Abonos recibidos de cartera", _money(v["abonos"])),
        ("Dinero real recibido", _money(v["recibido_real"])),
        ("Numero de ventas", _num(v["cantidad"])),
        ("Ticket promedio", _money(v["ticket_promedio"])),
    ]
    _add_table(doc, ["Concepto", "Valor"], venta_rows, col_widths_cm=[13.5, 3.5], right_cols=[1])

    _add_paragraph(
        doc,
        f"Comparativo con el periodo anterior ({cmp['previo']['inicio']} al {cmp['previo']['fin']}): "
        f"ventas {_num(cmp['crecimiento_facturado'])}% ({_money(cmp['facturado_anterior'])} anterior), "
        f"dinero recibido {_num(cmp['crecimiento_recibido'])}% y gastos {_num(cmp['crecimiento_gastos'])}%.",
        size=9.5,
        italic=True,
        color=GRAY,
    )

    if v["top_productos"]:
        _add_paragraph(doc, "Productos mas vendidos del periodo:", bold=True, size=11, space_after=4)
        top_rows = [
            [p["producto"], _num(p["unidades"]), _money(p["generado"]), _money(p["utilidad"])]
            for p in v["top_productos"]
        ]
        _add_table(
            doc,
            ["Producto", "Unidades", "Generado", "Utilidad estimada"],
            top_rows,
            col_widths_cm=[7.1, 2.3, 3.8, 3.8],
            right_cols=[1, 2, 3],
        )

    # ---------- CAJA Y LIQUIDEZ ----------
    _add_section_heading(doc, "4. CAJA Y DINERO DISPONIBLE")
    saldo_rows = []
    total_disponible = 0.0
    for metodo, saldo in data["caja"]["saldo_por_metodo"].items():
        label = {
            "efectivo": "Efectivo",
            "nequi": "Nequi",
            "bancolombia": "Bancolombia",
            "bogota": "Banco de Bogota",
            "credito": "Cartera (fiado, no disponible)",
        }.get(metodo, metodo)
        if metodo == "credito":
            saldo_rows.append([label, _money(saldo)])
        else:
            total_disponible += saldo
            color = GREEN if saldo >= 0 else RED
            saldo_rows.append([label, _money(saldo)])
    saldo_rows.append(["TOTAL DISPONIBLE EN CAJA", _money(total_disponible)])
    _add_table(doc, ["Medio", "Saldo"], saldo_rows, col_widths_cm=[13.5, 3.5], right_cols=[1], last_bold=True)

    # ---------- GASTOS ----------
    _add_section_heading(doc, "5. GASTOS")
    gasto_rows = [
        ("Total gastos", _money(g["total"])),
        ("Cantidad de gastos", _num(g["cantidad"])),
        ("Retiros o saques del negocio", _money(g["retiros"])),
    ]
    _add_table(doc, ["Concepto", "Valor"], gasto_rows, col_widths_cm=[13.5, 3.5], right_cols=[1])

    if g["por_categoria"]:
        _add_paragraph(doc, "Gastos por categoria:", bold=True, size=11, space_after=4)
        cat_rows = sorted(g["por_categoria"].items(), key=lambda kv: -kv[1])
        labels = {
            "general": "General", "arriendo": "Arriendo", "servicios": "Servicios",
            "nomina": "Nomina", "transporte": "Transporte", "material": "Material",
            "impuestos": "Impuestos", "marketing": "Marketing", "mantenimiento": "Mantenimiento",
            "otro": "Otro",
        }
        _add_table(
            doc,
            ["Categoria", "Valor"],
            [[labels.get(k, k), _money(val)] for k, val in cat_rows],
            col_widths_cm=[13.5, 3.5],
            right_cols=[1],
        )

    # ---------- CARTERA ----------
    _add_section_heading(doc, "6. CARTERA Y DEUDORES")
    car = data["cartera"]
    cartera_rows = [
        ("Total vendido a credito", _money(car["deuda_total"])),
        ("Abonos recibidos (todo el historial)", _money(car["abonos_total"])),
        ("Pendiente por cobrar", _money(car["deuda_pendiente"])),
        ("Deudores activos", _num(car["deudores_count"])),
    ]
    _add_table(doc, ["Concepto", "Valor"], cartera_rows, col_widths_cm=[13.5, 3.5], right_cols=[1])

    _add_paragraph(doc, "Antiguedad de la deuda:", bold=True, size=11, space_after=4)
    ant_rows = []
    for bucket in ["0-30", "31-60", "61-90", "91+"]:
        label = {"0-30": "0 a 30 dias", "31-60": "31 a 60 dias", "61-90": "61 a 90 dias", "91+": "Mas de 90 dias"}[bucket]
        monto = car["antiguedad"].get(bucket, 0)
        count = car["antiguedad_count"].get(bucket, 0)
        color = GREEN if bucket == "0-30" else (ORANGE if bucket in ("31-60", "61-90") else RED)
        ant_rows.append([label, _num(count), _money(monto)])
    _add_table(doc, ["Rango", "Deudores", "Saldo"], ant_rows, col_widths_cm=[8.0, 3.0, 6.0], right_cols=[1, 2])

    if car["deudores"]:
        _add_paragraph(doc, "Deudores detallados:", bold=True, size=11, space_after=4)
        deudor_rows = [
            [d["cliente"], d["remision"], d["fecha"], f"{d['dias']} dias", _money(d["saldo"])]
            for d in car["deudores"][:20]
        ]
        _add_table(
            doc,
            ["Cliente", "Remision", "Fecha", "Antiguedad", "Saldo"],
            deudor_rows,
            col_widths_cm=[5.5, 3.0, 2.5, 2.8, 3.2],
            right_cols=[1, 2, 3, 4],
        )
        if len(car["deudores"]) > 20:
            _add_paragraph(doc, f"... y {len(car['deudores']) - 20} deudores mas.", size=9, italic=True, color=GRAY)

    # ---------- RENTABILIDAD ----------
    _add_section_heading(doc, "7. RENTABILIDAD")
    inv = data["inventario"]
    rent_rows = [
        ("Dinero real recibido", _money(v["recibido_real"])),
        ("(-) Gastos del periodo", _money(g["total"])),
        ("Ganancia de caja", _money(r["ganancia_caja"])),
        ("Margen de caja", f"{_num(r['margen_caja'])}%"),
        ("(-) Costo estimado de la mercancia vendida", _money(r["costo_mercancia"])),
        ("Utilidad comercial estimada", _money(r["utilidad_comercial"])),
        ("Margen comercial", f"{_num(r['margen_comercial'])}%"),
        ("Utilidad potencial del inventario (venta - costo)", _money(inv["utilidad_potencial"])),
    ]
    _add_table(doc, ["Concepto", "Valor"], rent_rows, col_widths_cm=[13.5, 3.5], right_cols=[1])

    prest_colchon_rows = [
        ("Prestamos internos pendientes de cobro", f"{_money(data['prestamos']['pendiente'])}  ({_num(data['prestamos']['activos'])} activos)"),
        ("Colchon financiero: base objetivo", _money(data["colchon"]["base"])),
        ("Colchon financiero: ahorro acumulado", _money(data["colchon"]["ahorrado"])),
        ("Prestamos hechos desde el colchon, por cobrar", _money(data["colchon"]["pendiente_cobrar"])),
    ]
    _add_paragraph(doc, "Prestamos internos y colchon financiero:", bold=True, size=11, space_after=4)
    _add_table(doc, ["Concepto", "Valor"], prest_colchon_rows, col_widths_cm=[13.5, 3.5], right_cols=[1])

    # ---------- INVENTARIO ----------
    _add_section_heading(doc, "8. INVENTARIO")
    inv_rows = [
        ("Valor del inventario a costo de compra", _money(inv["costo"])),
        ("Valor del inventario a precio de venta", _money(inv["venta"])),
        ("Total productos activos", _num(inv["total_productos"])),
        ("Productos en o bajo el stock minimo", _num(inv["bajo_stock"])),
        ("Productos agotados (stock 0)", _num(inv["agotados"])),
        ("Productos sin margen (venta <= costo)", _num(inv["sin_margen"])),
    ]
    _add_table(doc, ["Concepto", "Valor"], inv_rows, col_widths_cm=[13.5, 3.5], right_cols=[1])

    # ---------- LO BUENO ----------
    _add_section_heading(doc, "9. LO BUENO DE LA EMPRESA")
    if data["fortalezas"]:
        for f in data["fortalezas"]:
            _add_bullet(doc, f["titulo"], f["detalle"], label_color=GREEN)
    else:
        _add_paragraph(doc, "Sin fortalezas destacadas en este periodo.", italic=True, color=GRAY)

    # ---------- LO CRITICO ----------
    _add_section_heading(doc, "10. PUNTOS CRITICOS A ATENDER")
    if data["alertas"]:
        severidad_color = {"CRITICA": RED, "ALTA": ORANGE, "MEDIA": GOLD_DARK, "BAJA": GRAY}
        for a in data["alertas"]:
            _add_bullet(
                doc,
                f"[{a['severidad']}] {a['titulo']}",
                a["detalle"],
                label_color=severidad_color.get(a["severidad"], GRAY),
            )
    else:
        _add_paragraph(doc, "Sin puntos criticos detectados. Excelente gestion.", italic=True, color=GREEN)

    # ---------- PROYECCION ----------
    _add_section_heading(doc, "11. PROYECCION PARA LOS PROXIMOS MESES")
    proy = data["proyeccion"]
    _add_paragraph(
        doc,
        f"Basado en {proy['periodo_dias']} dias observados y una tendencia de {_num(proy['tendencia'])}%, "
        f"el promedio de ingreso diario es {_money(proy['promedio_diario'])}.",
        size=10.5,
    )
    _add_paragraph(doc, proy["resumen"])

    esc = proy["escenarios"]
    esc_rows = [
        [
            "Pesimista (-10pp)",
            _money(esc["pesimista"]["ventas_mes"]),
            _money(esc["pesimista"]["ventas_3m"]),
            _money(esc["pesimista"]["utilidad_mes"]),
            RED,
        ],
        [
            "Realista (tendencia actual)",
            _money(esc["realista"]["ventas_mes"]),
            _money(esc["realista"]["ventas_3m"]),
            _money(esc["realista"]["utilidad_mes"]),
            GOLD_DARK,
        ],
        [
            "Optimista (+10pp)",
            _money(esc["optimista"]["ventas_mes"]),
            _money(esc["optimista"]["ventas_3m"]),
            _money(esc["optimista"]["utilidad_mes"]),
            GREEN,
        ],
    ]
    table = doc.add_table(rows=1, cols=4)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, h in enumerate(["Escenario", "Ventas / mes", "Ventas en 3 meses", "Utilidad / mes"]):
        _set_cell_shading(table.rows[0].cells[i], GOLD_DARK)
        _set_cell_text(table.rows[0].cells[i], h, bold=True, color="FFFFFF", align="center")
    for row in esc_rows:
        cells = table.add_row().cells
        for i in range(4):
            _set_cell_text(
                cells[i], row[i],
                bold=True,
                color=row[4] if i == 3 else DARK,
                align="right" if i > 0 else "left",
            )

    _add_paragraph(doc, "", space_after=2)
    _add_paragraph(doc, "Gastos proyectados mensuales:", bold=True, size=11, space_after=2)
    _add_paragraph(doc, f"{_money(proy['gastos_mensuales'])} (*proyeccion lineal del promedio diario del periodo).", size=10)

    _add_paragraph(doc, "Plan de accion recomendado:", bold=True, size=11, space_after=4)
    for rec in proy["recomendaciones"]:
        _add_bullet(doc, "", rec, label_color=GOLD_DARK)

    # ---------- CIERRE ----------
    _add_section_heading(doc, "12. CONCLUSION")
    conclusion = (
        "Este informe se genera automaticamente con los datos registrados en el sistema de gestion de JORMAR DISTRIBUCIONES. "
        "Sirve de hoja de ruta para tomar decisiones comerciales y financieras en los proximos meses. "
        "Para maximizar su utilidad, mantenga los registros de ventas, gastos, abonos y retiros al dia."
    )
    _add_paragraph(doc, conclusion, italic=True, color=GRAY)

    firma = doc.add_paragraph()
    firma.alignment = WD_ALIGN_PARAGRAPH.CENTER
    firma.paragraph_format.space_before = Pt(16)
    run_f = firma.add_run("JORMAR DISTRIBUCIONES\nDireccion Comercial y Financiera")
    run_f.font.name = "Calibri"
    run_f.font.size = Pt(9)
    run_f.font.bold = True
    run_f.font.color.rgb = RGBColor.from_string(GOLD_DARK)

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()